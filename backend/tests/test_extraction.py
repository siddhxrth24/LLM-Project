import tempfile
from pathlib import Path

import pytest

from app.services.extraction_service import (
    chunk_text,
    extract_pdf,
    extract_pptx,
    sanitize_filename,
)


def test_sanitize_filename() -> None:
    """Test path traversal prevention and character sanitization."""
    unsafe = "../../etc/passwd/some crazy*name?.pdf"
    clean = sanitize_filename(unsafe)
    assert ".." not in clean
    assert "/" not in clean
    assert "\\" not in clean
    assert "*" not in clean
    assert "?" not in clean
    assert clean.endswith(".pdf")

    empty_sanitized = sanitize_filename("")
    assert empty_sanitized == "unnamed_document"


def test_chunk_text_overlap_and_metadata() -> None:
    """Test text chunking with sliding window and metadata preservation."""
    # 300 words text
    words = [f"word{i}" for i in range(300)]
    text = " ".join(words)
    pages = [(1, text)]

    chunk_size = 100
    overlap = 20
    chunks = chunk_text(
        pages=pages,
        chunk_size=chunk_size,
        overlap=overlap,
        source_type="pdf",
        source_filename="lecture1.pdf",
    )

    assert len(chunks) > 1
    # Check first chunk
    assert chunks[0]["chunk_index"] == 0
    assert chunks[0]["page_number"] == 1
    assert chunks[0]["slide_number"] is None
    assert chunks[0]["source_filename"] == "lecture1.pdf"
    assert chunks[0]["metadata"]["source_type"] == "pdf"

    # Check overlap: end of chunk 0 should overlap start of chunk 1
    c0_words = chunks[0]["content"].split(" ")
    c1_words = chunks[1]["content"].split(" ")
    assert c0_words[-overlap:] == c1_words[:overlap]


def test_extract_pdf_in_memory() -> None:
    """Test PDF text extraction using a dynamically generated PyMuPDF PDF."""
    try:
        import fitz
    except ImportError:
        pytest.skip("PyMuPDF (fitz) not installed in current environment")

    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tmp:
        tmp_path = Path(tmp.name)

    try:
        doc = fitz.open()
        page1 = doc.new_page()
        page1.insert_text((50, 72), "Chapter 1: Process Synchronization and Deadlocks")
        page2 = doc.new_page()
        page2.insert_text((50, 72), "Chapter 2: Memory Hierarchy and Cache Mapping")
        doc.save(str(tmp_path))
        doc.close()

        extracted = extract_pdf(tmp_path)
        assert len(extracted) == 2
        assert extracted[0][0] == 1
        assert "Process Synchronization" in extracted[0][1]
        assert extracted[1][0] == 2
        assert "Memory Hierarchy" in extracted[1][1]
    finally:
        if tmp_path.exists():
            tmp_path.unlink()


def test_extract_pptx_in_memory() -> None:
    """Test PPTX extraction using a dynamically generated presentation."""
    try:
        from pptx import Presentation
        from pptx.util import Inches
    except ImportError:
        pytest.skip("python-pptx not installed in current environment")

    with tempfile.NamedTemporaryFile(suffix=".pptx", delete=False) as tmp:
        tmp_path = Path(tmp.name)

    try:
        prs = Presentation()
        # Slide 1
        blank_slide_layout = prs.slide_layouts[6]
        slide1 = prs.slides.add_slide(blank_slide_layout)
        txBox = slide1.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2))
        tf = txBox.text_frame
        tf.text = "Unit 1: Introduction to Computer Networks"

        # Slide 2 with table
        slide2 = prs.slides.add_slide(blank_slide_layout)
        table_shape = slide2.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(2))
        table = table_shape.table
        table.cell(0, 0).text = "OSI Layer"
        table.cell(0, 1).text = "Function"
        table.cell(1, 0).text = "Physical"
        table.cell(1, 1).text = "Bit Transmission"

        prs.save(str(tmp_path))

        extracted = extract_pptx(tmp_path)
        assert len(extracted) == 2
        assert extracted[0][0] == 1
        assert "Computer Networks" in extracted[0][1]
        assert extracted[1][0] == 2
        assert "OSI Layer" in extracted[1][1]
        assert "Bit Transmission" in extracted[1][1]
    finally:
        if tmp_path.exists():
            tmp_path.unlink()
